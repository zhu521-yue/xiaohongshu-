from __future__ import annotations

import csv
import json
from pathlib import Path

from .models import ExtractionResult, PageResult, TableResult
from .utils import ensure_directory


def write_outputs(result: ExtractionResult) -> None:
    output_dir = Path(result.output_dir)
    ensure_directory(output_dir)
    write_page_texts(result.pages, output_dir / "pages")
    write_tables(result.pages, output_dir / "tables")
    write_metadata(result, output_dir / "metadata.json")
    write_markdown(result, output_dir / "document.md")


def write_page_texts(pages: list[PageResult], pages_dir: Path) -> None:
    ensure_directory(pages_dir)
    for page in pages:
        path = pages_dir / f"page_{page.page_number:03d}.txt"
        path.write_text(page.text, encoding="utf-8")


def write_tables(pages: list[PageResult], tables_dir: Path) -> None:
    ensure_directory(tables_dir)
    tables: list[TableResult] = []
    for page in pages:
        for table in page.tables:
            path = tables_dir / f"page_{page.page_number:03d}_table_{table.index:02d}.csv"
            table.csv_path = f"tables/{path.name}"
            _write_csv(path, table.rows)
            tables.append(table)

    if tables:
        _write_xlsx(tables_dir / "tables.xlsx", tables)


def write_metadata(result: ExtractionResult, path: Path) -> None:
    data = result.to_dict(include_text=False)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def write_markdown(result: ExtractionResult, path: Path) -> None:
    lines: list[str] = []
    lines.append(f"# PDF Extraction: {Path(result.source_pdf).name}")
    lines.append("")
    lines.append(f"- Pages processed: {result.processed_pages}/{result.page_count}")
    lines.append(f"- Text characters: {result.text_length}")
    lines.append(f"- Tables: {result.table_count}")
    lines.append(f"- Figures: {result.figure_count}")
    lines.append("")

    for page in result.pages:
        lines.append(f"## Page {page.page_number}")
        lines.append("")
        if page.text.strip():
            lines.append(page.text.strip())
        else:
            lines.append("_No text extracted._")
        lines.append("")

        if page.tables:
            lines.append("### Tables")
            lines.append("")
            for table in page.tables:
                path_text = table.csv_path or ""
                lines.append(
                    f"- Table {table.index}: {table.row_count} rows x "
                    f"{table.column_count} columns"
                    + (f" (`{path_text}`)" if path_text else "")
                )
            lines.append("")

        if page.figures:
            lines.append("### Figures")
            lines.append("")
            for figure in page.figures:
                lines.append(f"- {figure.kind} {figure.index}: `{figure.path}`")
            lines.append("")

    path.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")


def _write_csv(path: Path, rows: list[list[str]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.writer(file)
        writer.writerows(rows)


def _write_xlsx(path: Path, tables: list[TableResult]) -> None:
    try:
        from openpyxl import Workbook
    except ImportError:
        return

    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    used_names: set[str] = set()
    for table in tables:
        title = f"p{table.page_number}_t{table.index}"
        title = _unique_sheet_name(title[:31], used_names)
        sheet = workbook.create_sheet(title)
        for row in table.rows:
            sheet.append(row)

    workbook.save(path)


def _unique_sheet_name(name: str, used_names: set[str]) -> str:
    candidate = name or "table"
    index = 1
    while candidate in used_names:
        suffix = f"_{index}"
        candidate = f"{name[:31 - len(suffix)]}{suffix}"
        index += 1
    used_names.add(candidate)
    return candidate
